#!/usr/bin/env python
"""Convert monthly FactResellerSales pipe-CSV exports to data-grade xlsx files.

Input : export_reseller_sales/csv/FactResellerSales_YYYY-MM.csv
Output: export_reseller_sales/FactResellerSales_YYYY-MM.xlsx

Values are typed (ints, floats, datetimes, text) so the workbooks are directly
usable for analysis; NULLs become empty cells. Layout is A1-header data table
(machine-friendly), with styled header row, frozen header, auto-fitted widths.
"""
import os
import sys

import pandas as pd
from openpyxl import load_workbook

XLSX_SKILL_DIR = r"C:\Users\phakk\.zcode\cli\plugins\cache\zcode-plugins-official\spreadsheets\0.2.1\skills\xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)
from base import fill_header, font_header, border_header, align_header, auto_fit_columns  # noqa: E402

CSV_DIR = os.path.join(os.path.dirname(__file__), "..", "export_reseller_sales", "csv")
OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "export_reseller_sales")

INT_COLS = ["ProductKey", "OrderDateKey", "DueDateKey", "ShipDateKey", "ResellerKey",
            "EmployeeKey", "PromotionKey", "CurrencyKey", "SalesTerritoryKey"]
NULLABLE_INT_COLS = ["SalesOrderLineNumber", "RevisionNumber", "OrderQuantity"]
FLOAT_COLS = ["UnitPrice", "ExtendedAmount", "UnitPriceDiscountPct", "DiscountAmount",
              "ProductStandardCost", "TotalProductCost", "SalesAmount", "TaxAmt", "Freight"]
TEXT_COLS = ["SalesOrderNumber", "CarrierTrackingNumber", "CustomerPONumber"]
DATE_COLS = ["OrderDate", "DueDate", "ShipDate"]
SHEET = "FactResellerSales"


def convert(csv_path: str, xlsx_path: str) -> int:
    df = pd.read_csv(csv_path, sep="|", dtype=str, na_values=["NULL", ""], keep_default_na=True)
    for c in INT_COLS:
        df[c] = df[c].astype("int64")
    for c in NULLABLE_INT_COLS:
        df[c] = df[c].astype("Int64")
    for c in FLOAT_COLS:
        df[c] = df[c].astype("float64")
    for c in DATE_COLS:
        df[c] = pd.to_datetime(df[c], errors="raise")

    df.to_excel(xlsx_path, index=False, sheet_name=SHEET, engine="openpyxl")

    wb = load_workbook(xlsx_path)
    ws = wb[SHEET]
    for cell in ws[1]:
        cell.fill = fill_header()
        cell.font = font_header()
        cell.alignment = align_header()
        cell.border = border_header()
    for col_idx, header in enumerate(df.columns, start=1):
        letter = ws.cell(row=1, column=col_idx).column_letter
        if header in DATE_COLS:
            for row in range(2, ws.max_row + 1):
                ws.cell(row=row, column=col_idx).number_format = "YYYY-MM-DD HH:MM:SS"
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    auto_fit_columns(ws, min_width=8, max_width=32, header_row=1, data_start_row=2)
    wb.properties.creator = "Z.ai"
    wb.save(xlsx_path)
    return len(df)


def main() -> None:
    files = sorted(f for f in os.listdir(CSV_DIR) if f.endswith(".csv"))
    total = 0
    for name in files:
        csv_path = os.path.join(CSV_DIR, name)
        xlsx_path = os.path.join(OUT_DIR, name.replace(".csv", ".xlsx"))
        rows = convert(csv_path, xlsx_path)
        total += rows
        print(f"{os.path.basename(xlsx_path)}  {rows} rows")
    print(f"TOTAL: {total} rows in {len(files)} xlsx files")


if __name__ == "__main__":
    main()
