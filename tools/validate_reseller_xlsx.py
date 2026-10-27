#!/usr/bin/env python
"""Validate monthly FactResellerSales xlsx files against SQL reference data.

Checks per file:
  1. data row count matches the CSV it was converted from
  2. SUM(SalesAmount) matches the SQL per-month sum (tolerance 0.01)
  3. every OrderDate value falls in the file's own month
"""
import math
import os
import sys

from openpyxl import load_workbook

BASE = os.path.join(os.path.dirname(__file__), "..", "export_reseller_sales")
XLSX_DIR = BASE
CSV_DIR = os.path.join(BASE, "csv")
ORDERDATE_COL = 25   # 1-based
SALESAMOUNT_COL = 20

sql_sums = {}
with open(os.path.join(BASE, "sql_monthly_sums.txt"), encoding="utf-8") as fh:
    for line in fh:
        line = line.strip()
        if "|" in line:
            ym, val = line.split("|")
            sql_sums[ym] = float(val)

errors = []
files = sorted(f for f in os.listdir(XLSX_DIR) if f.endswith(".xlsx"))
total_rows = 0
total_sum = 0.0
for name in files:
    ym = name.replace("FactResellerSales_", "").replace(".xlsx", "")
    csv_rows = sum(1 for _ in open(os.path.join(CSV_DIR, name.replace(".xlsx", ".csv")), encoding="utf-8")) - 1

    wb = load_workbook(os.path.join(XLSX_DIR, name), read_only=True)
    ws = wb["FactResellerSales"]
    rows = 0
    s = 0.0
    bad_month = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None and row[ORDERDATE_COL - 1] is None:
            continue
        rows += 1
        v = row[SALESAMOUNT_COL - 1]
        if v is not None:
            s += float(v)
        od = row[ORDERDATE_COL - 1]
        if od is not None and od.strftime("%Y-%m") != ym:
            bad_month += 1
    wb.close()

    total_rows += rows
    total_sum += s
    msgs = []
    if rows != csv_rows:
        msgs.append(f"row mismatch: xlsx={rows} csv={csv_rows}")
    if ym in sql_sums and abs(s - sql_sums[ym]) > 0.01:
        msgs.append(f"sum mismatch: xlsx={s:.4f} sql={sql_sums[ym]:.4f}")
    if bad_month:
        msgs.append(f"{bad_month} rows outside month {ym}")
    if msgs:
        errors.append(f"{name}: " + "; ".join(msgs))
        print(f"FAIL {name}: {'; '.join(msgs)}")
    else:
        print(f"PASS {name}: {rows} rows, SalesAmount={s:.4f}")

print("---")
print(f"TOTAL: {total_rows} rows, SalesAmount={total_sum:.4f} (SQL=96309252.3782)")
if total_rows != 61267 or abs(total_sum - 96309252.3782) > 0.01:
    errors.append("grand total mismatch")
print("RESULT:", "ALL PASS" if not errors else f"{len(errors)} FAILURES")
sys.exit(1 if errors else 0)
