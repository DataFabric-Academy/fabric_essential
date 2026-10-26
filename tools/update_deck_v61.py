# -*- coding: utf-8 -*-
"""ปรับเด็ค MS-FB-101 จาก Version 6 เป็น Version 6.1:
1) แก้ข้อความที่ล้าสมัย (default semantic model, สะกดผิด)
2) ใส่ speaker notes "v6.1:" ให้ผู้สอน
3) เพิ่มภาคผนวกท้ายเด็ค
รันจากราก repo:  python -X utf8 tools/update_deck_v61.py
"""
import copy, re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

SRC = "Trainer_Docs/9Expert Training-MS-FB-101-Aj-Phakkapong-Version 6.pptx"
DST = "Trainer Docs".replace(" ", "_")  # unused placeholder
DST = "Trainer_Docs/9Expert Training-MS-FB-101-Aj-Phakkapong-Version 6.1.pptx"

prs = Presentation(SRC)
log = {"para": 0, "notes": 0}

def iter_text_frames(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from iter_text_frames(sh.shapes)
            continue
        if sh.has_text_frame:
            yield sh.text_frame
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    yield cell.text_frame

def para_text(p):
    return "".join(r.text for r in p.runs)

def set_para(p, text):
    if not p.runs:
        p.add_run().text = text
        return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""

# ---------- 1) แก้ข้อความ ----------
def apply_para_rules(p):
    t = para_text(p)
    if not t.strip():
        return
    orig = t
    rules = [
        # (เงื่อนไขย่อยหรือ None=match ตรง, old_substring, new_substring)
        (None, "Power BI semantic model (Default) ถูกสร้างขึ้น",
               "Semantic model สร้างเองด้วยปุ่ม New semantic model เมื่อ Tables และ Views พร้อม (default model ยกเลิกแล้ว — ภาคผนวก v6.1)"),
        (None, "Power BI semantic model (Default)",
               "Semantic model สร้างเองด้วย New semantic model (default model ยกเลิกแล้ว — ภาคผนวก v6.1)"),
        (None, "Default ถูกสร้างอัตโนมัติพร้อม Warehouse",
               "สร้างด้วยปุ่ม New semantic model (default model ยกเลิกแล้ว — ภาคผนวก v6.1)"),
        (None, "Default ถูกสร้างอัตโนมัติพร้อม Lakehouse",
               "สร้างด้วยปุ่ม New semantic model จาก SQL analytics endpoint (default model ยกเลิกแล้ว — ภาคผนวก v6.1)"),
        (None, "Manage default Power BI semantic Model",
               "สร้าง semantic model จาก SQL analytics endpoint ด้วยปุ่ม New semantic model (ปุ่ม Manage default model ถูกถอดออกแล้ว — ภาคผนวก v6.1)"),
        (None, "Get date", "Get data"),
        (None, "DimGeoghapy", "DimGeography"),
        (None, "Combile files", "Combine files"),
        (None, "อัฟโหลด", "อัปโหลด"),
        (None, "Version 6 Publish Date: 1 July 2026",
               "Version 6.1 — Update: 25 October 2026 (แก้ไขจาก Version 6 — ดูภาคผนวกท้ายเล่ม)"),
        # หัวข้อสไลด์ชุด "แก้ไข Semantic Model (default)"
        ("Semantic Model", "(default)", "(สร้างเอง)"),
    ]
    new_t = t
    for cond, old, new in rules:
        if cond is None:
            new_t = new_t.replace(old, new)
        else:
            if cond in new_t and old in new_t:
                new_t = new_t.replace(old, new)
                new_t = re.sub(r"\s{2,}", " ", new_t)
    if new_t != orig:
        set_para(p, new_t)
        log["para"] += 1

for slide in prs.slides:
    for tf in iter_text_frames(slide.shapes):
        for p in tf.paragraphs:
            apply_para_rules(p)

# ---------- 2) speaker notes ----------
NOTES = {
    "สร้าง Warehouse แรก":
        "v6.1: Default semantic model ถูกยกเลิกทั้งกระบวนการ (ตั้งแต่ปลายปี 2025) — ระบบไม่สร้างให้อัตโนมัติอีก ให้ห้องใช้ปุ่ม New semantic model แทน (โมดูล 07) ประกาศทางการ: community.fabric.microsoft.com/blog/.../sunsetting-default-semantic-models",
    "Dataflow GEN 2":
        "v6.1: ของใหม่ที่ควรเสริม — Fast Copy (ingest ข้อมูลใหญ่บาง connector วิ่งบน Copy engine) และตัวแปร Dataflow Gen2 with CI/CD สำหรับเก็บ definition ลง Git (Git-integrated ยังเรียกจาก pipeline ไม่ได้ช่วง preview)",
    "สรุปการนำเข้า":
        "v6.1: มีตัวใหม่ Copy job — low-code สำหรับงาน copy ล้วน มี incremental/CDC ในตัว เรียกจาก pipeline ได้ผ่าน Copy job activity เลือกยังไง: copy ล้วน→Copy job, transform→Dataflow, ผสมหลาย activity→Pipeline",
    "อาจเคยรู้จักในชื่อ":
        "v6.1: Default semantic model ถูกยกเลิกแล้ว — สอนแนวสร้างเองด้วย New semantic model (โมดูล 07)",
    "Runtime 1.2":
        "v6.1: Runtime 2.0 GA (Spark 4.1, Delta 4) และกำลังเป็น default ของ item ใหม่; Runtime 1.3 เข้า LTS ถึง มี.ค. 2027; Runtime 1.1/1.2 หมดการสนับสนุน — ตรวจรุ่นจริงบน tenant ก่อนสอน (learn.microsoft.com/fabric/data-engineering/runtime-2-0)",
    "SQL analytics endpoint  เพื่อเป็นส่วนติดต่อ":
        "v6.1: Default semantic model ยกเลิกแล้ว — ใช้ New semantic model จาก Lakehouse/SQL endpoint",
    "Semantic Model บน SQL":
        "v6.1: สร้าง semantic model จาก SQL analytics endpoint ด้วยปุ่ม New semantic model (default model ไม่มีแล้ว)",
    "Manage default Power BI semantic Model":
        "v6.1: ปุ่ม Manage default Power BI semantic model ถูกถอดออกพร้อมการ sunset default models — ขั้นตอนนี้กลายเป็น: ที่ SQL analytics endpoint กด New semantic model เลือกตาราง แล้วจัดการใน web model editor",
    "เลือกตารางดังแสดง":
        "v6.1: ทำบน semantic model ที่สร้างเอง (New semantic model) ไม่ใช่ default",
    "และในอนาคต Fabric Semantic Model":
        "v6.1: ทำบน semantic model ที่สร้างเอง (New semantic model) ไม่ใช่ default",
    "ISSUE: Mark as date table":
        "v6.1: นอกจากสมมติฐานเรื่อง type/ค่าซ้ำ — ให้ทดสอบบน semantic model ที่สร้างใหม่เสมอ (default model ไม่มีแล้ว) เฉลยเต็มที่ solutions/15-lh-semantic-model (repo)",
    "On-demand load":
        "v6.1: คอลัมน์ default semantic model ในตารางเทียบ — ปัจจุบันสร้างเองด้วย New semantic model",
    "เพจ Premium":
        "v6.1: ยืนยันแล้ว 60 วัน / 64 CU / OneLake 1 TB — เพิ่มเติม: tenant ใช้ trial ได้สูงสุด 5 ครั้ง และ item ถูกลบ 7 วันหลัง trial หมดอายุ (learn.microsoft.com/fabric/fundamentals/fabric-trial)",
    "License ประเภท F":
        "v6.1: ตรวจราคา/SKU ปัจจุบันที่ learn.microsoft.com/fabric/enterprise/licenses ก่อนคาบ",
    "เข้าใจ Fabric Capacity":
        "v6.1: Copilot ใช้ได้กับ paid Fabric SKUs (F2 ขึ้นไป) ไม่รวม Pro/PPU — เกี่ยวข้องเมื่อเทียบ SKU",
    "เข้าใจ Power BI Capacity":
        "v6.1: Copilot ใช้ได้กับ paid Fabric SKUs (F2 ขึ้นไป) ไม่รวม Pro/PPU — ตรวจราคาปัจจุบันก่อนคาบ",
}

for slide in prs.slides:
    all_text = "\n".join(p.text for tf in iter_text_frames(slide.shapes) for p in tf.paragraphs)
    for marker, note in NOTES.items():
        if marker in all_text:
            slide.notes_slide.notes_text_frame.text = note
            log["notes"] += 1
            break

# ---------- 3) ภาคผนวกท้ายเด็ค ----------
def blank_layout(prs):
    best = None
    for layout in prs.slide_layouts:
        if len(layout.placeholders) == 0:
            return layout
        if best is None or len(layout.placeholders) < len(best.placeholders):
            best = layout
    return best

def add_appendix_slide(prs, title, bullets):
    layout = blank_layout(prs)
    slide = prs.slides.add_slide(layout)
    # ล้าง placeholder ที่ติดมากับ layout (ถ้ามี)
    for sh in list(slide.shapes):
        el = sh._element
        el.getparent().remove(el)
    W, H = prs.slide_width, prs.slide_height
    dark = RGBColor(0x20, 0x20, 0x20)
    gray = RGBColor(0x80, 0x80, 0x80)

    tb = slide.shapes.add_textbox(Inches(0.45), H * 0.06, W - Inches(0.9), H * 0.14)
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = dark

    bb = slide.shapes.add_textbox(Inches(0.55), H * 0.24, W - Inches(1.1), H * 0.66)
    tf = bb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = "•  " + line
        para.font.size = Pt(17)
        para.font.color.rgb = dark
        para.space_after = Pt(8)

    fb = slide.shapes.add_textbox(Inches(0.55), H * 0.93, W - Inches(1.1), H * 0.05)
    fp = fb.text_frame.paragraphs[0]
    fr = fp.add_run(); fr.text = "Copyright © 2026 9Expert Co., Ltd. All rights reserved. | MS-FB-101 v6.1 Update — October 2026"
    fr.font.size = Pt(10); fr.font.color.rgb = gray
    return slide

APPENDIX = [
    ("ภาคผนวก v6.1 — อัปเดตให้ทันสมัย (ตุลาคม 2026)", [
        "อัปเดตตาม MS Learn และประกาศทางการของ Microsoft ณ วันที่ 25 ต.ค. 2026",
        "หัวใจใหญ่ที่สุด: Default semantic models ถูกยกเลิกทั้งกระบวนการ",
        "ของใหม่ที่แนะนำในหลักสูตร: Copy job, Runtime 2.0, Fast Copy, OneLake security",
        "เอกสารประกอบ (repo) ไม่อ้างเลขหน้าสไลด์ — อ้างด้วยชื่อหัวข้อ/item/ปุ่ม เพื่อให้เด็คแก้-เรียงใหม่ได้อิสระ",
    ]),
    ("Default Semantic Model ถูกยกเลิกแล้ว", [
        "ก.ย. 2025 — ระบบไม่สร้าง default model อัตโนมัติอีก (Warehouse / Lakehouse / SQL DB / Mirrored items)",
        "ปลาย ต.ค. 2025 — default models เดิมถูกแปลงเป็น regular semantic models",
        "ปลาย ธ.ค. 2025 — ปิดรอบพร้อมถอด Reporting tab ที่พึ่ง default model",
        "สิ่งที่ต้องทำ: สร้าง model ด้วยปุ่ม New semantic model เสมอ — ออกแบบเฉพาะคำถามธุรกิจ",
        "รายงานเดิมที่ต่อ model ที่ถูกแปลงแล้วยังใช้งานได้ปกติ",
    ]),
    ("วิธีทำงานกับ Semantic Model วันนี้", [
        "จาก Warehouse / Lakehouse / SQL analytics endpoint → ปุ่ม New semantic model",
        "เลือกเฉพาะตาราง/วิวที่จำเป็น — model เล็กที่ตรงคำถามดีกว่าใหญ่ที่คลุมเครือ",
        "จัดการ relationship, hierarchy, hide, sort ใน web model editor",
        "อ้างอิง: learn.microsoft.com/fabric/data-warehouse/semantic-models",
    ]),
    ("Fabric Trial ปัจจุบัน", [
        "60 วัน / 64 capacity units / OneLake สูงสุด 1 TB",
        "tenant ละ 5 trial — บัญชีที่เคยใช้แล้วเปิดซ้ำไม่ได้",
        "7 วันหลังหมดอายุ: item ทั้งหมดถูกลบ",
        "เปิดก่อนวันอบรมเสมอ — อย่ารอวันจริง",
    ]),
    ("Spark Runtime วันนี้", [
        "Runtime 2.0 GA — Spark 4.1, Delta 4 กำลังเป็น default ของ item ใหม่",
        "Runtime 1.3 — เข้า Long Term Support จนถึง มี.ค. 2027 (Spark 3.5.5, Python 3.11)",
        "Runtime 1.1 / 1.2 — หมดการสนับสนุน",
        "ตรวจรุ่นที่ tenant จริงก่อนสอน + อ่าน migration guide ก่อนย้าย 1.3 → 2.0",
    ]),
    ("Dataflow Gen2 วันนี้", [
        "Fast Copy — ingest ข้อมูลใหญ่บาง connector วิ่งบน Copy engine โดยตรง",
        "Dataflow Gen2 with CI/CD — เก็บ definition ลง Git (ช่วง preview)",
        "ข้อจำกัดปัจจุบัน: Git-integrated dataflow ยังเรียกจาก pipeline ไม่ได้",
        "เลือกปลายทาง + Replace ยังเป็นหัวใจเดิมของ Gen2",
    ]),
    ("Copy job — ตัวใหม่ใน Data Factory", [
        "งาน copy ล้วน low-code — สร้างเร็วกว่า pipeline",
        "incremental copy + CDC ในตัว (Azure SQL / SQL MI / SQL Server GA)",
        "เรียกจาก pipeline ได้ผ่าน Copy job activity",
        "เลือกยังไง: copy ล้วน → Copy job, transform → Dataflow, ผสมหลาย activity → Pipeline",
    ]),
    ("Lakehouse วันนี้", [
        "schema-enabled lakehouse — เปิด schema ตอนสร้าง (มี dbo ให้ default)",
        "SQL analytics endpoint สะท้อน schema ปัจจุบันของ Delta tables เสมอ",
        "Materialized Lake Views — view ที่ persist และ refresh เอง",
        "หลัก Tables/Files + managed/external ยังใช้เหมือนเดิม",
    ]),
    ("Copilot ใน Fabric", [
        "ใช้ได้กับ paid Fabric SKUs (F2 ขึ้นไป) — ไม่รวม Pro/PPU",
        "Warehouse: natural language → SQL, code completion, quick actions",
        "Notebook: สร้าง PySpark/Python code; Data Factory: ช่วยเขียน pipeline/dataflow",
        "input สูงสุด 10,000 ตัวอักษร — ตรวจเงื่อนไข tenant ก่อนเปิดใช้ในห้อง",
    ]),
    ("OneLake วันนี้", [
        "OneLake security GA (พ.ค. 2026) — สิทธิ์ระดับโฟลเดอร์/ตาราง + row/column-level security",
        "Shortcut transformations — แปลงข้อมูลระหว่างทางตอน shortcut ดึงเข้ามา",
        "Shortcut ยังเป็น pointer ไม่ copy — หัวใจของ OneLake คงเดิม",
    ]),
    ("ศัพท์ที่ควรใช้ / เลี่ยง", [
        "ใช้: \"semantic model (สร้างเอง)\" — เลี่ยง: \"default model\" (ยกเลิกแล้ว)",
        "ใช้: \"SQL analytics endpoint\" — เลี่ยง: \"SQL Analytic Endpoint\" (สะกดเก่า)",
        "ใช้: \"Copy job\" สำหรับงาน copy ล้วน — ไม่ยึดว่าต้อง pipeline เท่านั้น",
        "ใช้: Runtime 2.0 / 1.3 — เลี่ยงอ้าง runtime เก่า (1.1 / 1.2)",
    ]),
    ("ผู้สอนต้องเดินตรวจก่อนคาบ", [
        "Trial ทุกบัญชี (เปิดล่วงหน้า) + บัญชีสำรองอย่างน้อย 2",
        "Blob/SQL credentials จาก Link.txt — connect จริงหนึ่งรอบ",
        "Tenant: default model ไม่มีแล้ว / runtime ที่ใช้ได้ / Copilot เปิดหรือไม่",
        "NYC demo: โหลดล่วงหน้า + ใช้ Warehouse แยก",
        "รายการเต็ม: instructor/preclass-checklist.md (repo)",
    ]),
    ("แหล่งอ้างอิง", [
        "Sunsetting default semantic models — community.fabric.microsoft.com (ประกาศทางการ)",
        "Fabric trial — learn.microsoft.com/fabric/fundamentals/fabric-trial",
        "Runtime 2.0 — learn.microsoft.com/fabric/data-engineering/runtime-2-0",
        "Copy job — learn.microsoft.com/fabric/data-factory/what-is-copy-job",
        "Dataflow Gen2 CI/CD — learn.microsoft.com/fabric/data-factory/dataflow-gen2-cicd-and-git-integration",
        "Copilot in Fabric — learn.microsoft.com/fabric/fundamentals/copilot-fabric-overview",
        "แผนที่เต็ม: reference/ms-learn-map.md (repo)",
    ]),
]
for title, bullets in APPENDIX:
    add_appendix_slide(prs, title, bullets)

prs.save(DST)
print(f"saved: {DST}")
print(f"paragraphs edited: {log['para']}, notes added: {log['notes']}, total slides now: {len(prs.slides)}")
